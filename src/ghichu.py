"""
Xuat lai file NGUON GOC kem chu thich - moi tab giu nguyen nhu ban dau.

Khac han file trong results/: file kia la danh sach da lam sach, gop het 12 tab
thanh mot bang. File nay giu dung hinh dang file goc ben cung cap gui - tung
tab, tung dong, tung cot y nguyen - chi them 3 cot chu thich o cuoi moi dong:

    Ket luan | Ma loi | Chu thich

de nguoi doc mo file quen thuoc cua minh len la thay ngay tung link bi gi.
Chi to mau phan chu thich, khong dong vao du lieu goc.

Chu thich khong chi lay tu ket qua check. Nhung link bi loai ngay tu buoc lam
sach (trung lap, tro ve chinh money site, o khong phai URL...) cung duoc ghi ro
ly do - day chinh la nhung dong bien mat khoi backlinks_master.csv ma nhin file
ket qua thuong khong hieu vi sao.

Chay:
    python src/cli.py ghichu -c config/checkbacklink.yaml
"""

import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).parent))
import bl_config
import diagnose as D
import doisoat
import ingest as ingest_mod
import urlutil as U

# --------------------------------------------------------------------- mau sac
# Nhom mau chia theo cau hoi "link nay co dung duoc khong", khong phai theo
# muc do nghiem trong - de nguoi doc luot mot lan la phan loai duoc ca tab.
XANH = PatternFill("solid", fgColor="C6EFCE")   # dung duoc
DO = PatternFill("solid", fgColor="FFC7CE")     # mat / khong dung duoc
VANG = PatternFill("solid", fgColor="FFEB9C")   # chua ket luan duoc
CAM = PatternFill("solid", fgColor="FFD9B3")    # bi loai tu buoc lam sach
XAM = PatternFill("solid", fgColor="EDEDED")    # khong lien quan / chua check

F_XANH = Font(color="006100", bold=True)
F_DO = Font(color="9C0006", bold=True)
F_VANG = Font(color="9C6500", bold=True)
F_CAM = Font(color="974706", bold=True)
F_XAM = Font(color="808080")

HDR_FILL = PatternFill("solid", fgColor="305496")
HDR_FONT = Font(color="FFFFFF", bold=True)

# nhom -> (fill, font, nhan ngan hien o cot "Ket luan")
NHOM = {
    "SONG":        (XANH, F_XANH, "Link con"),
    "MAT":         (DO,   F_DO,   "Link mat"),
    "CHECK_TAY":   (VANG, F_VANG, "Phai check tay"),
    "TRUNG":       (CAM,  F_CAM,  "Trung lap"),
    "MONEY":       (CAM,  F_CAM,  "Link cua minh"),
    "URL_HONG":    (CAM,  F_CAM,  "URL hong"),
    "DOMAIN_LOAI": (CAM,  F_CAM,  "Domain bi loai"),
    "CHUA_CHECK":  (XAM,  F_XAM,  "Chua check"),
    "KHONG_TINH":  (XAM,  F_XAM,  "Khong tinh"),
}

COT_THEM = [("Ket luan", 16), ("Ma loi", 22), ("Chu thich", 78)]


def _cot_cuoi_co_data(cells):
    for j in range(len(cells) - 1, -1, -1):
        if str(cells[j] if cells[j] is not None else "").strip():
            return j + 1
    return 0


class _Ghi:
    """Mot chu thich cho mot URL trong o."""

    __slots__ = ("nhom", "code", "text")

    def __init__(self, nhom, code, text):
        self.nhom, self.code, self.text = nhom, code, text


# ------------------------------------------------------------------ phan loai
def _ghi_tu_ket_qua(k):
    """Bien mot dong ket qua check thanh chu thich."""
    v = (k.get("ket_luan") or "").strip() or D.V_CHECK
    code = (k.get("diag_code") or "").strip()
    phan = [k.get("chan_doan") or ""]
    # Tool ket luan theo Googlebot. Ghi ca hai goc nhin de nguoi doc hieu vi
    # sao mot link "mo ra van thay" lai bi tinh la mat - va nguoc lai.
    nx = (k.get("nguoi_xem") or "").strip()
    gb = (k.get("googlebot") or "").strip()
    if nx and gb:
        phan.append("Nguoi xem thay: %s / Google thay: %s" % (nx, gb))
    tren = (k.get("dich_tang_tren") or "").strip()
    if tren.startswith("DA CHET"):
        phan.append("URL tang tren %s: %s" % (tren, k.get("points_to") or ""))
    viec = (k.get("viec_can_lam") or "").strip()
    if viec:
        yc = (k.get("yeu_cau") or "").strip()
        phan.append("Viec can lam%s: %s" % (" (%s)" % yc if yc else "", viec))
    them = (k.get("canh_bao_them") or "").strip()
    if them:
        phan.append("Canh bao them: " + them)
    http = (k.get("http_code") or "").strip()
    if http and http not in ("0", ""):
        phan.append("HTTP " + http)
    return _Ghi(v if v in NHOM else "CHECK_TAY", code,
                " | ".join(p for p in phan if p))


def _phan_loai_url(u_raw, cfg, kq, seen, tier, sheet_name):
    """Mot URL tho -> chu thich. Lap lai dung thu tu loc cua ingest.build()
    de ly do ghi ra khop chinh xac voi ly do link bi bo khoi master CSV."""
    minlen = cfg.ingest["min_url_length"]
    strip_tr = cfg.ingest["strip_tracking"]
    own = set([cfg.money_domain] + list(cfg.extra_domains))
    drop_domains = [d.lower() for d in cfg.ingest["drop_domains"]]

    u = U.clean_raw(u_raw)
    if not u or not U.is_valid(u, minlen):
        return _Ghi("URL_HONG", "URL_HONG",
                    "O nay khong phai URL mo duoc nen khong tinh la link da giao.")

    dom = U.domain_of(u)
    if dom in own or U.registrable(dom) in own:
        return _Ghi("MONEY", "TRO_VE_MONEY_SITE",
                    "URL tro ve chinh money site - day la link cua minh, "
                    "khong phai backlink ben ngoai.")
    if any(d in dom for d in drop_domains):
        return _Ghi("DOMAIN_LOAI", "DOMAIN_BI_LOAI",
                    "Domain nam trong ingest.drop_domains - khong bao gio dung "
                    "lam backlink.")

    key = U.normalize(u, strip_tr)
    if cfg.ingest["dedupe"] and key in seen:
        g_tier, g_tab = seen[key]
        cung = " (trung ngay trong cung tab nay)" if g_tab == sheet_name else ""
        return _Ghi("TRUNG", "LINK_TRUNG",
                    "URL nay da xuat hien truoc do o tab '%s' (tier %s)%s. "
                    "Chi tinh duoc mot lan." % (g_tab, g_tier, cung))
    seen[key] = (tier, sheet_name)

    k = kq.get(U.normalize(u))
    if k:
        return _ghi_tu_ket_qua(k)
    return _Ghi("CHUA_CHECK", "",
                "Chua co ket qua check cho URL nay (chua chay den, hoac lan chay "
                "truoc dung --limit).")


def _urls_trong_o(cell):
    """Cac URL nam trong mot o. O khong co URL nao thi tra ve list rong."""
    s = str(cell or "").strip()
    if not s:
        return []
    found = U.extract_urls(s)
    if found:
        return found
    one = U.clean_raw(s)
    return [one] if one and U.is_valid(one, 8) else []


# ---------------------------------------------------------------------- build
def build(cfg, url="", thu_muc=None):
    """Doc nguon goc + ket qua check moi nhat -> du lieu de ghi ra xlsx.

    Tra ve (danh_sach_tab, file_ket_qua_da_dung, thong_ke).
    Moi tab: {"ten", "tier", "bo_qua", "rows": [(cells, ghi_chu_list)], ...}
    """
    groups = ingest_mod.doc_nguon(cfg, url, giu_kieu=True)
    kq, file_kq = doisoat.doc_ket_qua(cfg, thu_muc)
    drop_sheets = [s.lower() for s in cfg.ingest["drop_sheets"]]

    seen = {}
    dem = Counter()
    tabs = []

    for sheet_name, sheet_rows in groups:
        bo_qua = ""
        tier = None
        if any(d in sheet_name.lower() for d in drop_sheets):
            bo_qua = ("Tab bi loai theo ingest.drop_sheets - khong nap, "
                      "khong check.")
        else:
            tier = cfg.tier_for_sheet(sheet_name)
            if tier is None:
                m = ingest_mod.TIER_HEADING_RE.search(sheet_name)
                if m and int(m.group(1)) in cfg.tiers:
                    tier = int(m.group(1))
            if tier is None:
                bo_qua = ("Ten tab khong khop tu khoa 'match' cua tier nao - "
                          "toan bo tab bi bo qua. Them tu khoa vao config.")

        # Tab co khai bao sheet_columns thi chi cot do moi tinh la backlink.
        want = ingest_mod._column_rule(cfg, sheet_name)
        cot_tinh, dong_tieu_de, ten_cot = None, None, []
        if want and not bo_qua:
            dong_tieu_de, idx, ten_cot = ingest_mod.tim_cot(sheet_rows, want)
            if idx:
                cot_tinh = set(idx)

        rows_out = []
        for i, row in enumerate(sheet_rows):
            ghis = []
            if not bo_qua:
                for j, cell in enumerate(row):
                    if cot_tinh is not None and i <= (dong_tieu_de or 0):
                        continue
                    for u_raw in _urls_trong_o(cell):
                        if cot_tinh is not None and j not in cot_tinh:
                            ghis.append(_Ghi(
                                "KHONG_TINH", "COT_KHONG_TINH",
                                "Cot nay khong duoc khai trong ingest.sheet_columns "
                                "(chi lay cot: %s) nen khong tinh la cho dat "
                                "backlink." % ", ".join(ten_cot)))
                            continue
                        ghis.append(_phan_loai_url(
                            u_raw, cfg, kq, seen, tier, sheet_name))
            for g in ghis:
                dem[g.nhom] += 1
            rows_out.append((list(row), ghis))

        tabs.append({"ten": sheet_name, "tier": tier, "bo_qua": bo_qua,
                     "rows": rows_out, "ten_cot": ten_cot})

    return tabs, file_kq, dem


# ------------------------------------------------------------------ ghi ra file
def _gop(ghis):
    """Nhieu URL tren cung mot dong -> mot bo (nhom, ket_luan, ma_loi, chu_thich).

    Nhom cua ca dong lay theo cai xau nhat, de mau canh bao khong bi mot link
    tot cung dong che mat.
    """
    uu_tien = ["MAT", "TRUNG", "MONEY", "URL_HONG", "DOMAIN_LOAI",
               "CHECK_TAY", "CHUA_CHECK", "KHONG_TINH", "SONG"]
    # O nam ngoai cot da khai trong sheet_columns (vd cot "Ten Mien") chi la
    # ghi chu phu. Co link that tren cung dong thi de link that quyet dinh
    # nhan va mau cua ca dong.
    chinh = [g for g in ghis if g.nhom != "KHONG_TINH"] or ghis
    phu = [g for g in ghis if g.nhom == "KHONG_TINH"] if chinh is not ghis else []

    nhom = min((g.nhom for g in chinh),
               key=lambda n: uu_tien.index(n) if n in uu_tien else 99)
    xep = chinh + phu
    codes = ", ".join(sorted({g.code for g in xep if g.code}))

    if len(chinh) == 1:
        g = chinh[0]
        text = g.text
        if phu:
            text += "  •  " + phu[0].text
        return nhom, NHOM[g.nhom][2], codes, text

    dem = Counter(NHOM[g.nhom][2] for g in chinh)
    nhan = "%d link: %s" % (len(chinh),
                            ", ".join("%d %s" % (n, t) for t, n in dem.most_common()))
    text = "  •  ".join("%s: %s" % (NHOM[g.nhom][2], g.text) for g in xep)
    return nhom, nhan, codes, text


def write_xlsx(tabs, path, cfg, file_kq=(), dem=None):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)

    _sheet_huong_dan(wb, cfg, file_kq, dem or Counter())

    dung_ten = set()
    for tab in tabs:
        ten = _ten_hop_le(tab["ten"], dung_ten)
        ws = wb.create_sheet(ten)

        # Google Sheet hay tra ve mot loat cot rong o cuoi - cat bot de cot
        # chu thich nam ngay canh du lieu that, khong troi ra xa.
        n_cot_goc = max((_cot_cuoi_co_data(r[0]) for r in tab["rows"]), default=1) or 1
        c0 = n_cot_goc + 2          # chua mot cot trong lam vach ngan

        for cells, ghis in tab["rows"]:
            ws.append(list(cells))
            if not ghis:
                continue
            nhom, nhan, code, text = _gop(ghis)
            fill, font, _ = NHOM[nhom]
            rn = ws.max_row
            for off, val in enumerate((nhan, code, text)):
                c = ws.cell(row=rn, column=c0 + off, value=val)
                c.fill = fill
                c.font = font if off == 0 else Font(color=font.color.rgb)
                c.alignment = Alignment(vertical="top", wrap_text=(off == 2))

        # tieu de cho 3 cot them, dat o dong 1 de luon nhin thay
        for off, (ten_cot, w) in enumerate(COT_THEM):
            c = ws.cell(row=1, column=c0 + off, value=ten_cot)
            c.fill, c.font = HDR_FILL, HDR_FONT
            ws.column_dimensions[get_column_letter(c0 + off)].width = w

        for j in range(1, n_cot_goc + 1):
            ws.column_dimensions[get_column_letter(j)].width = 42
        ws.freeze_panes = "A2"

        if tab["bo_qua"]:
            c = ws.cell(row=2, column=c0, value=tab["bo_qua"])
            c.fill, c.font = XAM, F_XAM
            c.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(path)
    return path


def _ten_hop_le(ten, dung):
    """Excel: ten sheet toi da 31 ky tu, khong duoc chua : \\ / ? * [ ]"""
    s = "".join("-" if ch in ':\\/?*[]' else ch for ch in str(ten))[:31] or "Tab"
    goc, i = s, 2
    while s.lower() in dung:
        hau = "_%d" % i
        s, i = goc[:31 - len(hau)] + hau, i + 1
    dung.add(s.lower())
    return s


def _sheet_huong_dan(wb, cfg, file_kq, dem):
    ws = wb.create_sheet("Doc truoc")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 96

    ws.append(["FILE GOC CO CHU THICH"])
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    ws.append(["Giu nguyen tung tab, tung dong, tung cot cua file ben cung cap "
               "gui. Ba cot mau o ben phai la phan tool them vao."])
    ws.cell(row=2, column=1).font = Font(italic=True, color="7F7F7F")
    ws.append([])
    ws.append(["Money site", "", cfg.money_domain])
    ws.append(["Ket qua check dung", "", ", ".join(file_kq) if file_kq
               else "CHUA CHAY CHECK - moi link deu ghi 'Chua check'"])
    ws.append([])

    ws.append(["Nhan", "So link", "Nghia la gi"])
    for c in range(1, 4):
        ws.cell(row=ws.max_row, column=c).fill = HDR_FILL
        ws.cell(row=ws.max_row, column=c).font = HDR_FONT

    y_nghia = [
        ("SONG", "Tool doc duoc trang va nhin thay the <a>. Link dung duoc."),
        ("MAT", "Link khong con: 404, 410, bai bi go, domain het han, hoac "
                "vao duoc trang ma khong con link ve dich." if cfg.loose else
                "Link khong con gia tri: 404, 410, bai bi go, domain het han, "
                "hoac trang mang the noindex."),
        ("CHECK_TAY", "Tool chua doc duoc noi dung that (chan bot, captcha, "
                      "tuong dang nhap, chua render JS). CHUA phai la link mat."),
        ("TRUNG", "URL da xuat hien o dong/tab truoc. Chi tinh duoc mot lan."),
        ("MONEY", "URL tro ve chinh money site - khong phai backlink."),
        ("URL_HONG", "O du lieu khong phai URL mo duoc."),
        ("DOMAIN_LOAI", "Domain nam trong ingest.drop_domains."),
        ("KHONG_TINH", "Cot khong duoc khai trong ingest.sheet_columns."),
        ("CHUA_CHECK", "Chua chay check den URL nay."),
    ]
    for nhom, mo_ta in y_nghia:
        fill, font, nhan = NHOM[nhom]
        ws.append([nhan, dem.get(nhom, 0), mo_ta])
        rn = ws.max_row
        for c in range(1, 4):
            ws.cell(row=rn, column=c).fill = fill
            ws.cell(row=rn, column=c).font = font if c == 1 else Font(color=font.color.rgb)
        ws.cell(row=rn, column=3).alignment = Alignment(wrap_text=True, vertical="top")

    ws.append([])
    ws.append(["", "", "Mot dong co nhieu URL thi cot 'Ket luan' ghi tong hop va "
                       "to mau theo link xau nhat; cot 'Chu thich' liet ke tung link."])
    ws.cell(row=ws.max_row, column=3).font = Font(italic=True, color="7F7F7F")
    ws.freeze_panes = "A8"
    return ws

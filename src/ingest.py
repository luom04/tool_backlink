"""
Tai nguon backlink tu Google Sheets / Google Docs / file local,
lam sach va gop thanh mot master CSV chuan: stt,tier,sheet,source_url.

Chay:
    python src/ingest.py -c config/checkbacklink.yaml
    python src/ingest.py -c config/site.yaml --url "https://docs.google.com/..."
    python src/ingest.py -c config/site.yaml --dry-run

Yeu cau voi Google Sheets / Docs: file phai duoc chia se o che do
"Bat ky ai co duong lien ket" (Anyone with the link - Viewer).
"""

import argparse
import csv
import io
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import httpx

sys.path.insert(0, str(Path(__file__).parent))
import bl_config
import urlutil as U

SHEET_ID_RE = re.compile(r"/spreadsheets/d/([a-zA-Z0-9-_]+)")
DOC_ID_RE = re.compile(r"/document/d/([a-zA-Z0-9-_]+)")
TIER_HEADING_RE = re.compile(r"(?:tier|tang|tầng)\s*[:\-]?\s*([1-9])", re.I)


def _fetch(url: str, timeout: int = 60) -> bytes:
    with httpx.Client(follow_redirects=True, timeout=timeout, verify=False) as c:
        r = c.get(url)
        if r.status_code == 403 or "accounts.google.com" in str(r.url):
            raise SystemExit(
                "Google tu choi tai file (403 hoac bi day sang trang dang nhap).\n"
                "Mo file > Share > General access > 'Anyone with the link' > Viewer, "
                "roi chay lai."
            )
        r.raise_for_status()
        return r.content


def detect_source(url: str) -> str:
    if SHEET_ID_RE.search(url or ""):
        return "google_sheet"
    if DOC_ID_RE.search(url or ""):
        return "google_doc"
    if (url or "").lower().endswith((".csv", ".tsv")):
        return "csv"
    if (url or "").lower().endswith((".xlsx", ".xlsm")):
        return "xlsx"
    return "csv"


# ---------------------------------------------------------------- readers
def _tai_google_sheet(url: str) -> bytes:
    sid = SHEET_ID_RE.search(url).group(1)
    export = "https://docs.google.com/spreadsheets/d/%s/export?format=xlsx" % sid
    print("Tai Google Sheet %s ..." % sid, file=sys.stderr)
    return _fetch(export)


def read_google_sheet(url: str, giu_kieu: bool = False):
    """Tra ve list (ten_tab, rows) cho moi tab. rows la list cac dong,
    moi dong la list gia tri o - giu nguyen cot de con chon cot duoc."""
    return read_xlsx_bytes(_tai_google_sheet(url), giu_kieu)


def read_xlsx_bytes(blob: bytes, giu_kieu: bool = False):
    """giu_kieu=True: giu nguyen kieu du lieu goc (so van la so, ngay van la
    ngay) - dung khi can dung lai file goc nguyen ven. Mac dinh ep ve chuoi
    cho khau lam sach."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        if giu_kieu:
            rows = [list(row) for row in ws.iter_rows(values_only=True)]
        else:
            rows = [["" if v is None else str(v) for v in row]
                    for row in ws.iter_rows(values_only=True)]
        out.append((ws.title, rows))
    wb.close()
    return out


def read_google_doc(url: str):
    """Google Docs: tach nhom theo dong tieu de co chu 'Tier N' / 'Tang N'."""
    did = DOC_ID_RE.search(url).group(1)
    export = "https://docs.google.com/document/d/%s/export?format=txt" % did
    print("Tai Google Doc %s ..." % did, file=sys.stderr)
    text = _fetch(export).decode("utf-8", "replace")
    return _split_text_by_heading(text)


def _split_text_by_heading(text: str):
    groups, current, buf = [], "(khong ro nhom)", []
    for line in text.splitlines():
        s = line.strip()
        if not s:
            continue
        is_heading = not U.URL_RE.search(s) and len(s) < 120 and len(s.split()) <= 14
        if is_heading:
            if buf:
                groups.append((current, buf))
                buf = []
            current = s
            continue
        buf.append([s])
    if buf:
        groups.append((current, buf))
    return groups


def read_local(path: str, giu_kieu: bool = False):
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        return read_xlsx_bytes(p.read_bytes(), giu_kieu)
    with open(p, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    header = [h.lower().strip() for h in (rows[0] if rows else [])]
    if "source_url" in header and "sheet" in header:
        i_sheet, i_url = header.index("sheet"), header.index("source_url")
        by = defaultdict(list)
        for r in rows[1:]:
            if len(r) > max(i_sheet, i_url):
                by[r[i_sheet]].append([r[i_url]])
        return list(by.items())
    return [(p.stem, rows)]


def _column_rule(cfg, sheet_name):
    """Tra ve list ten cot can lay URL, neu tab nay co khai bao trong config."""
    rules = cfg.ingest.get("sheet_columns") or {}
    best, best_len = None, 0
    for pattern, cols in rules.items():
        pat = bl_config.fold(pattern)
        if pat and pat in bl_config.fold(sheet_name) and len(pat) > best_len:
            best, best_len = cols, len(pat)
    if best is None:
        return None
    return [bl_config.fold(c) for c in (best if isinstance(best, list) else [best])]


def tim_cot(rows, want_cols):
    """Do dong tieu de trong 10 dong dau -> (so_dong_tieu_de, chi_so_cot, ten_cot).

    Khong tim thay tieu de nao khop thi tra ve (None, [], []).

    Tab khong co dong tieu de thi khai theo chu cai cot: "cot D", "cot AB".
    Khi do lay tu dong dau tien (so_dong_tieu_de = -1) - o tieu de / STT khong
    phai URL nen tu bi loc o buoc sau.
    """
    chu = [re.fullmatch(r"cot\s*([a-z]{1,2})", w) for w in want_cols]
    if want_cols and all(chu):
        idx = []
        for m in chu:
            n = 0
            for ch in m.group(1):
                n = n * 26 + (ord(ch) - 96)
            idx.append(n - 1)
        return -1, idx, ["cot %s" % m.group(1).upper() for m in chu]
    for i, row in enumerate(rows[:10]):
        idx, names = [], []
        for j, v in enumerate(row):
            f = bl_config.fold(v)
            if f and any(f == w or w in f for w in want_cols):
                idx.append(j)
                names.append(str(v).strip())
        if idx:
            return i, idx, names
    return None, [], []


def _cells_from_columns(rows, want_cols, sheet_name):
    """Tim dong tieu de trong 10 dong dau, roi chi lay o thuoc cac cot da chon.

    Tra ve (danh_sach_o, ten_cot_tim_duoc). Neu khong tim thay tieu de nao khop
    thi tra ve (None, []) de goi y quay lai cach quet toan bo o.
    """
    i, idx, names = tim_cot(rows, want_cols)
    if i is None:
        return None, []
    cells = []
    for r in rows[i + 1:]:
        for j in idx:
            if j < len(r) and str(r[j]).strip():
                cells.append(str(r[j]))
    return cells, names


# ---------------------------------------------------------------- pipeline
class Stats(Counter):
    """Counter thuong, nhung gan them duoc so ke toan theo tung tab."""
    per_sheet = []


def _so_moi(ten, tier, bo_qua=""):
    """Mot dong so ke toan cho mot tab nguon."""
    return {"tab": ten, "tier": tier, "bo_qua": bo_qua,
            "tho": 0, "url_hong": 0, "money_site": 0, "domain_loai": 0,
            "trung": 0, "trung_trong_tab": 0, "trung_voi_tab_khac": 0,
            "trung_voi": Counter(), "nhan": 0, "domain_rieng": 0}


def trang_chu_thua(urls):
    """URL trong MOT dong -> tap URL la trang chu cua mot link khac cung dong.

    File ben cung cap hay co dang: | Ten mien | DA | Link dat |, tuc cung dong
    vua co trang chu (thong tin "dat tren site nao") vua co backlink that cung
    ten mien. Trang chu do khong phai backlink. Trang chu dung MOT MINH thi
    giu - co the backlink dat ngay trang chu (site ve tinh *.mystrikingly.com).

    "Cung ten mien" = cung host, hoac link nam tren ten mien con cua trang chu
    (loda-lang.org -> boinc.loda-lang.org). KHONG so theo ten mien goc: hai site
    a.mystrikingly.com va b.mystrikingly.com la hai ve tinh khac nhau.
    """
    from urllib.parse import urlparse
    info = []
    for raw in urls:
        u = U.clean_raw(raw) or ""
        p = urlparse(u if "://" in u else "https://" + u)
        info.append((raw, U.domain_of(u if "://" in u else "https://" + u),
                     p.path in ("", "/") and not p.query))
    bo = set()
    for raw, host, la_trang_chu in info:
        if not la_trang_chu or not host:
            continue
        if any(r2 != raw and not tc2 and (h2 == host or h2.endswith("." + host))
               for r2, h2, tc2 in info):
            bo.add(raw)
    return bo


def doc_nguon(cfg, override_url: str = "", giu_kieu: bool = False):
    """Doc nguon GOC ve dung nguyen trang: [(ten_tab, cac_dong)].

    Chua loc gi ca - giu nguyen tung tab, tung dong, tung o. build() dung ham
    nay roi moi lam sach; ghichu.py dung chinh no de dung lai file goc.
    """
    src_type = cfg.source_type
    url = override_url or cfg.source_url
    if override_url or src_type in ("auto", "", None):
        src_type = detect_source(url) if url else "local"

    if src_type == "google_sheet":
        return read_google_sheet(url, giu_kieu)
    if src_type == "google_doc":
        return read_google_doc(url)
    if url and src_type == "xlsx":
        return read_xlsx_bytes(_fetch(url), giu_kieu)
    if url and src_type == "csv":
        text = _fetch(url).decode("utf-8", "replace")
        return [("import", list(csv.reader(io.StringIO(text))))]
    return read_local(cfg.source_file or cfg.master_csv, giu_kieu)


def build(cfg, override_url: str = "", verbose: bool = True):
    groups = doc_nguon(cfg, override_url)

    drop_sheets = [s.lower() for s in cfg.ingest["drop_sheets"]]
    drop_domains = [d.lower() for d in cfg.ingest["drop_domains"]]
    strip_tr = cfg.ingest["strip_tracking"]
    minlen = cfg.ingest["min_url_length"]
    own = set([cfg.money_domain] + list(cfg.extra_domains))

    seen, rows = {}, []
    stats = Stats()
    unmatched_sheets = []
    # key da chuan hoa -> ho so mot URL bi lap:
    #   {"url": URL ban goc, "goc": (tier, tab), "lap": [(tier, tab, url_thay), ...]}
    dup_detail = {}

    # So ke toan theo tung tab nguon, de doi soat lai voi ben cung cap link.
    # Moi tab giu day du: tho -> tru tung khoan -> con lai thuc nhan.
    per_sheet = {}

    for sheet_name, sheet_rows in groups:
        if any(d in sheet_name.lower() for d in drop_sheets):
            stats["sheet_bo_qua"] += 1
            per_sheet[sheet_name] = _so_moi(sheet_name, None, "Tab bi loai theo drop_sheets")
            if verbose:
                print("  bo qua tab '%s' (ingest.drop_sheets)" % sheet_name, file=sys.stderr)
            continue

        tier = cfg.tier_for_sheet(sheet_name)
        if tier is None:
            m = TIER_HEADING_RE.search(sheet_name)
            if m and int(m.group(1)) in cfg.tiers:
                tier = int(m.group(1))
        if tier is None:
            unmatched_sheets.append(sheet_name)
            per_sheet[sheet_name] = _so_moi(sheet_name, None, "Khong khop tier nao")
            continue

        want = _column_rule(cfg, sheet_name)
        cells, picked = (None, [])
        if want:
            cells, picked = _cells_from_columns(sheet_rows, want, sheet_name)
            if cells is None:
                print("  CANH BAO: tab '%s' co khai bao sheet_columns nhung khong "
                      "tim thay dong tieu de khop -> quet toan bo o." % sheet_name,
                      file=sys.stderr)
        # Gom URL theo TUNG DONG de luat trang_chu_thua() so duoc cac o cung
        # dong voi nhau. Tab da chon cot thi moi o la mot nhom rieng.
        if cells is None:
            nhom_o = [[c for c in r if str(c).strip()] for r in sheet_rows]
        else:
            nhom_o = [[c] for c in cells]

        so = per_sheet.setdefault(sheet_name, _so_moi(sheet_name, tier))
        n_sheet = 0
        for o_cung_dong in nhom_o:
            found = []
            for cell in o_cung_dong:
                f = U.extract_urls(cell)
                if not f:
                    one = U.clean_raw(cell)
                    f = [one] if one else []
                found += f
            bo = trang_chu_thua(found) if cells is None else set()
            for raw in found:
                if raw in bo:
                    stats["trang_chu_cung_dong"] += 1
                    continue
                so["tho"] += 1
                u = U.clean_raw(raw)
                if not u or not U.is_valid(u, minlen):
                    stats["url_hong"] += 1
                    so["url_hong"] += 1
                    continue
                dom = U.domain_of(u)
                if dom in own or U.registrable(dom) in own:
                    stats["tu_tro_ve_money_site"] += 1
                    so["money_site"] += 1
                    continue
                if any(d in dom for d in drop_domains):
                    stats["domain_bi_loai"] += 1
                    so["domain_loai"] += 1
                    continue
                key = U.normalize(u, strip_tr)
                if cfg.ingest["dedupe"] and key in seen:
                    stats["trung_lap"] += 1
                    so["trung"] += 1
                    g_tier, g_tab, g_url = seen[key]
                    if g_tab == sheet_name:
                        so["trung_trong_tab"] += 1
                    else:
                        so["trung_voi_tab_khac"] += 1
                        so["trung_voi"][g_tab] += 1
                    ho_so = dup_detail.setdefault(
                        key, {"url": g_url, "goc": (g_tier, g_tab), "lap": []})
                    ho_so["lap"].append((tier, sheet_name, u))
                    continue
                seen[key] = (tier, sheet_name, u)
                rows.append({"tier": tier, "sheet": sheet_name, "source_url": u})
                so["nhan"] += 1
                n_sheet += 1
        so["domain_rieng"] = len({U.registrable(r["source_url"]) for r in rows
                                  if r["sheet"] == sheet_name})
        if verbose:
            extra = "  [chi lay cot: %s]" % ", ".join(picked) if picked else ""
            print("  tab '%s' -> tier %s: %d link%s"
                  % (sheet_name, tier, n_sheet, extra), file=sys.stderr)
        stats["tier_%s" % tier] += n_sheet

    rows.sort(key=lambda r: (r["tier"], r["sheet"]))
    for i, r in enumerate(rows, 1):
        r["stt"] = i
    stats.per_sheet = list(per_sheet.values())
    return rows, stats, unmatched_sheets, dup_detail


def write_master(rows, path):
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    with open(p, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["stt", "tier", "sheet", "source_url"])
        w.writeheader()
        for r in rows:
            w.writerow({k: r[k] for k in ("stt", "tier", "sheet", "source_url")})
    return p


def main():
    ap = argparse.ArgumentParser(
        description="Gop + lam sach nguon backlink thanh master CSV")
    ap.add_argument("-c", "--config", required=True)
    ap.add_argument("--url", default="", help="ghi de source.url trong config")
    ap.add_argument("-o", "--output", default="", help="ghi de source.master_csv")
    ap.add_argument("--dry-run", action="store_true", help="chi in thong ke, khong ghi file")
    args = ap.parse_args()

    cfg = bl_config.load(args.config)
    rows, stats, unmatched, dups = build(cfg, args.url)

    print("\n" + "=" * 58, file=sys.stderr)
    print("Tong link sach : %d" % len(rows), file=sys.stderr)
    for t in sorted(cfg.tiers):
        n = sum(1 for r in rows if r["tier"] == t)
        print("  tier %d (%s): %d" % (t, cfg.tiers[t].label, n), file=sys.stderr)
    for k in ("trung_lap", "url_hong", "tu_tro_ve_money_site", "domain_bi_loai",
              "trang_chu_cung_dong"):
        if stats[k]:
            print("  da loai [%s]: %d" % (k, stats[k]), file=sys.stderr)
    if cfg.ingest.get("warn_homepage_urls", True):
        from urllib.parse import urlparse
        bare = [r for r in rows
                if not (urlparse(r["source_url"]).path or "/").strip("/")
                and not urlparse(r["source_url"]).query]
        if bare:
            print("\nCANH BAO - %d URL chi co ten mien, khong co duong dan:"
                  % len(bare), file=sys.stderr)
            for r in bare[:12]:
                print("  tier %s | %s | %s" % (r["tier"], r["sheet"], r["source_url"]),
                      file=sys.stderr)
            if len(bare) > 12:
                print("  ... con %d dong nua" % (len(bare) - 12), file=sys.stderr)
            print("  Trang chu hiem khi la noi dat backlink. Rat co the cot do la"
                  " cot ten mien / DR chu khong phai cot link that.\n"
                  "  Khai bao 'ingest.sheet_columns' de chi lay dung cot link.",
                  file=sys.stderr)

    if unmatched:
        print("\nCANH BAO - tab/nhom khong khop tier nao (bi bo qua):", file=sys.stderr)
        for s in unmatched:
            print("  - %s" % s, file=sys.stderr)
        print("  Them tu khoa vao 'tiers.<N>.match' trong config de nhan dien.",
              file=sys.stderr)

    if args.dry_run:
        print("\n(dry-run: khong ghi file)", file=sys.stderr)
        return
    out = write_master(rows, args.output or cfg.master_csv)
    print("\nDa ghi: %s" % out, file=sys.stderr)


if __name__ == "__main__":
    main()

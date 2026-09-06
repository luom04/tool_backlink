"""
Xuat ket qua ra file XLSX co to mau theo muc do nghiem trong.

- Sheet "Chi tiet"      : tung link, to mau ca dong theo cot muc_do.
- Sheet "Tong hop"      : khoi CHOT LAI (con / mat / phai check tay) roi den
                          bang theo tier va theo ma loi.
- Sheet "Can check tay" : rieng nhung link tool chua ket luan duoc, kem cach xu ly.
- Dong dau duoc dong bang (freeze) va bat auto-filter de loc nhanh.
"""

import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).parent))
import diagnose as D

# mau nen theo muc do (1 nang nhat -> 5 tot)
FILL = {
    1: PatternFill("solid", fgColor="FFC7CE"),   # do nhat - CHET
    2: PatternFill("solid", fgColor="FFD9B3"),   # cam    - NANG
    3: PatternFill("solid", fgColor="FFF2CC"),   # vang   - CANH BAO
    4: PatternFill("solid", fgColor="DDEBF7"),   # xanh nhat - GHI CHU
    5: PatternFill("solid", fgColor="E2EFDA"),   # xanh la - TOT
}
FONT = {
    1: Font(color="9C0006", bold=True),
    2: Font(color="974706"),
    3: Font(color="7F6000"),
    4: Font(color="1F4E79"),
    5: Font(color="375623"),
}
HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(color="FFFFFF", bold=True)

COLUMNS = [
    ("stt", "STT", 6),
    ("tier", "Tier", 6),
    ("sheet", "Nhom nguon", 26),
    ("source_url", "URL nguon", 58),
    ("status", "Trang thai", 12),
    ("ket_luan", "Ket luan", 15),
    ("muc_do", "Muc do", 11),
    ("diag_code", "Ma loi", 24),
    ("chan_doan", "Chan doan", 58),
    ("viec_can_lam", "Viec can lam", 52),
    ("canh_bao_them", "Canh bao them", 34),
    ("http_code", "HTTP", 7),
    ("rel", "rel", 12),
    ("indexable", "Index?", 8),
    ("points_to", "Tro ve", 44),
    ("khop_tang", "Khop tang", 11),
    ("anchor_text", "Anchor text", 34),
    ("final_url", "URL cuoi", 44),
    ("rendered", "Cach doc", 11),
    ("elapsed", "Giay", 7),
    ("note", "Ghi chu ky thuat", 30),
    ("checked_at", "Thoi diem check", 17),
]


# mau rieng cho cot "Ket luan" - doc luot la thay ngay 3 nhom
V_FILL = {
    D.V_SONG:  PatternFill("solid", fgColor="C6EFCE"),
    D.V_MAT:   PatternFill("solid", fgColor="FFC7CE"),
    D.V_CHECK: PatternFill("solid", fgColor="FFEB9C"),
}
V_FONT = {
    D.V_SONG:  Font(color="006100", bold=True),
    D.V_MAT:   Font(color="9C0006", bold=True),
    D.V_CHECK: Font(color="9C6500", bold=True),
}


def _sev(r):
    try:
        return int(r.severity)
    except (TypeError, ValueError):
        return 3


def _verdict(r):
    return getattr(r, "ket_luan", "") or D.V_CHECK


def write_xlsx(results, path, cfg=None):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # ------------------------------------------------------- sheet chi tiet
    ws = wb.active
    ws.title = "Chi tiet"
    ws.append([c[1] for c in COLUMNS])
    for i, (_, _, w) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 26

    # link hong len dau, trong cung muc do thi tier nho len truoc
    ordered = sorted(results, key=lambda r: (_sev(r), int(r.tier) if str(r.tier).isdigit() else 99,
                                             int(r.stt) if str(r.stt).isdigit() else 0))
    for r in ordered:
        sev = _sev(r)
        row = [getattr(r, key, "") for key, _, _ in COLUMNS]
        ws.append(row)
        rn = ws.max_row
        for i in range(1, len(COLUMNS) + 1):
            c = ws.cell(row=rn, column=i)
            c.fill = FILL[sev]
            c.font = FONT[sev]
            c.alignment = Alignment(vertical="top", wrap_text=(COLUMNS[i - 1][0] in
                                                               ("chan_doan", "viec_can_lam")))
        keys = [c[0] for c in COLUMNS]
        url_col = keys.index("source_url") + 1
        cell = ws.cell(row=rn, column=url_col)
        if r.source_url and len(r.source_url) < 250:
            cell.hyperlink = r.source_url
            cell.font = Font(color="0563C1", underline="single")

        v = _verdict(r)
        vc = ws.cell(row=rn, column=keys.index("ket_luan") + 1)
        vc.value = D.V_LABEL.get(v, v)
        vc.fill, vc.font = V_FILL[v], V_FONT[v]

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(COLUMNS)), ws.max_row)

    # ------------------------------------------------------- sheet tong hop
    s = wb.create_sheet("Tong hop")
    if cfg is not None:
        s.append(["Site", getattr(cfg, "site_name", "")])
        s.append(["Money site", getattr(cfg, "money_domain", "")])
        s.append([])

    # ---- khoi CHOT LAI: 3 con so nguoi dung can nhat, dat len tren cung
    d = D.dem_ket_luan(results)
    tong = d["TONG"] or 1
    s.append(["CHOT LAI - doc dong nay truoc"])
    s.cell(row=s.max_row, column=1).font = Font(bold=True, size=13)
    s.append(["Ket luan", "So link", "Ty le", "Nghia la gi", "Phai lam gi"])
    for c in range(1, 6):
        s.cell(row=s.max_row, column=c).fill = HEADER_FILL
        s.cell(row=s.max_row, column=c).font = HEADER_FONT

    chot = [
        (D.V_SONG, d[D.V_SONG],
         "Tool doc duoc trang va nhin thay the <a> tro ve money site.",
         "KHONG can mo tay. Trong so nay co %d link hoan hao (muc TOT)."
         % d["HOAN_HAO"]),
        (D.V_MAT, d[D.V_MAT],
         "Tool doc duoc trang va chac chan link khong con gia tri (404/410/bai "
         "bi go/trang noindex).",
         "KHONG can mo tay. Thay bang nguon moi."),
        (D.V_CHECK, d[D.V_CHECK],
         "Tool KHONG doc duoc noi dung that (chan bot, captcha, dang nhap, "
         "chua render JS...). Day CHUA phai ket luan mat link.",
         "PHAI xu ly: %d link chi can chay lai tool, %d link phai mo trinh duyet. "
         "Danh sach o sheet 'Can check tay'." % (d[D.X_TOOL], d[D.X_NGUOI])),
    ]
    for v, n, nghia, lam in chot:
        s.append([D.V_LABEL[v], n, "%.1f%%" % (100.0 * n / tong), nghia, lam])
        for c in range(1, 6):
            s.cell(row=s.max_row, column=c).fill = V_FILL[v]
            s.cell(row=s.max_row, column=c).font = V_FONT[v]
        s.cell(row=s.max_row, column=4).alignment = Alignment(wrap_text=True, vertical="top")
        s.cell(row=s.max_row, column=5).alignment = Alignment(wrap_text=True, vertical="top")
        s.row_dimensions[s.max_row].height = 46
    s.append(["TONG", d["TONG"], "100.0%", "", ""])
    s.cell(row=s.max_row, column=1).font = Font(bold=True)
    s.cell(row=s.max_row, column=2).font = Font(bold=True)
    s.append([])

    s.append(["THEO TIER"])
    s.cell(row=s.max_row, column=1).font = Font(bold=True)
    s.append(["Tier", "Tong", "Link con", "Link mat", "Phai check tay",
              "CHET", "NANG"])
    for c in range(1, 8):
        s.cell(row=s.max_row, column=c).fill = HEADER_FILL
        s.cell(row=s.max_row, column=c).font = HEADER_FONT

    per = defaultdict(lambda: Counter())
    for r in results:
        per[str(r.tier)][_verdict(r)] += 1
        per[str(r.tier)][r.muc_do] += 1
        per[str(r.tier)]["tong"] += 1
    for t in sorted(per, key=lambda x: int(x) if x.isdigit() else 99):
        c = per[t]
        s.append([t, c["tong"], c[D.V_SONG], c[D.V_MAT], c[D.V_CHECK],
                  c["CHET"], c["NANG"]])
        for col, v in ((3, D.V_SONG), (4, D.V_MAT), (5, D.V_CHECK)):
            if c[v]:
                s.cell(row=s.max_row, column=col).fill = V_FILL[v]
                s.cell(row=s.max_row, column=col).font = V_FONT[v]

    s.append([])
    s.append(["THEO MA LOI - sap theo muc do nghiem trong"])
    s.cell(row=s.max_row, column=1).font = Font(bold=True)
    s.append(["Ma loi", "Muc do", "So luong", "Chan doan", "Viec can lam"])
    for c in range(1, 6):
        s.cell(row=s.max_row, column=c).fill = HEADER_FILL
        s.cell(row=s.max_row, column=c).font = HEADER_FONT

    codes = Counter(r.diag_code for r in results if r.diag_code)
    rows = []
    for code, n in codes.items():
        base = D.CATALOG.get(code)
        if not base:
            continue
        sev, why, todo = base
        real = min((_sev(r) for r in results if r.diag_code == code), default=sev)
        rows.append((real, code, n, why, todo))
    for real, code, n, why, todo in sorted(rows, key=lambda x: (x[0], -x[2])):
        s.append([code, D.SEV_LABEL[real], n, why, todo])
        for c in range(1, 6):
            s.cell(row=s.max_row, column=c).fill = FILL[real]
            s.cell(row=s.max_row, column=c).font = FONT[real]

    for col, w in zip("ABCDE", (26, 11, 12, 62, 62)):
        s.column_dimensions[col].width = w

    # --------------------------------------------- sheet "Can check tay"
    # Danh sach cu the tung link tool CHUA ket luan duoc, kem cach xu ly.
    # Nhom "Chay lai tool" len dau vi lam mot lan la xong ca cum.
    ck = wb.create_sheet("Can check tay")
    todo = [r for r in results if _verdict(r) == D.V_CHECK]

    ck.append(["DANH SACH LINK TOOL CHUA KET LUAN DUOC"])
    ck.cell(row=1, column=1).font = Font(bold=True, size=13)
    ck.append(["Nhung link duoi day KHONG phai la link mat. Tool bi chan hoac "
               "chua doc duoc noi dung that nen khong dam ket luan."])
    ck.cell(row=2, column=1).font = Font(italic=True, color="7F7F7F")
    ck.append([])

    CK_COLS = [("Cach xu ly", 16), ("Tier", 6), ("Nhom nguon", 24),
               ("URL nguon - bam de mo", 62), ("Ma loi", 22),
               ("Vi sao chua ket luan duoc", 56), ("Lam gi de biet chac", 58)]
    hdr = ck.max_row + 1
    ck.append([c[0] for c in CK_COLS])
    for i, (_, w) in enumerate(CK_COLS, 1):
        ck.column_dimensions[get_column_letter(i)].width = w
        cell = ck.cell(row=hdr, column=i)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ck.row_dimensions[hdr].height = 26

    if not todo:
        ck.append(["", "", "", "Khong co link nao phai check tay. "
                   "Toan bo ket qua deu da chac chan.", "", "", ""])
        ck.cell(row=ck.max_row, column=4).font = Font(color="006100", bold=True)
    else:
        # X_TOOL truoc X_NGUOI, trong moi nhom gom theo ma loi roi theo tier
        order = {D.X_TOOL: 0, D.X_NGUOI: 1}
        todo.sort(key=lambda r: (order.get(getattr(r, "cach_xu_ly", ""), 9),
                                 getattr(r, "cach_xu_ly", ""),
                                 r.diag_code,
                                 int(r.tier) if str(r.tier).isdigit() else 99))
        for r in todo:
            how = getattr(r, "cach_xu_ly", "") or D.X_NGUOI
            ck.append([how, r.tier, r.sheet, r.source_url, r.diag_code,
                       r.chan_doan,
                       getattr(r, "huong_dan_check", "") or r.viec_can_lam])
            rn = ck.max_row
            for i in range(1, len(CK_COLS) + 1):
                c = ck.cell(row=rn, column=i)
                c.fill = V_FILL[D.V_CHECK]
                c.font = V_FONT[D.V_CHECK] if i == 1 else Font(color="9C6500")
                c.alignment = Alignment(vertical="top", wrap_text=(i >= 6))
            if r.source_url and len(r.source_url) < 250:
                u = ck.cell(row=rn, column=4)
                u.hyperlink = r.source_url
                u.font = Font(color="0563C1", underline="single")
        ck.freeze_panes = "A%d" % (hdr + 1)
        ck.auto_filter.ref = "A%d:%s%d" % (hdr, get_column_letter(len(CK_COLS)),
                                           ck.max_row)

    # ------------------------------------------------------- chu giai mau
    g = wb.create_sheet("Chu giai")
    g.append(["KET LUAN - co phai mo tay kiem tra khong?"])
    g.cell(row=1, column=1).font = Font(bold=True, size=12)
    g.append(["Ket luan", "Mau", "Y nghia"])
    for c in range(1, 4):
        g.cell(row=g.max_row, column=c).fill = HEADER_FILL
        g.cell(row=g.max_row, column=c).font = HEADER_FONT
    for v, mean in (
        (D.V_SONG,  "Tool nhin thay the <a> va trang co the duoc index. Chac chan "
                    "con link. Khong can mo tay. Gom ca link tro sai tang."),
        (D.V_MAT,   "Tool doc duoc trang va chac chan link khong con gia tri: 404, "
                    "410, bai bi go, hoac trang mang the noindex (Google khong "
                    "doc toi nen the <a> vo nghia). Khong can mo tay."),
        (D.V_CHECK, "Tool chua doc duoc noi dung that. CHUA ket luan. Xem sheet "
                    "'Can check tay'."),
    ):
        g.append([D.V_LABEL[v], "", mean])
        for c in range(1, 4):
            g.cell(row=g.max_row, column=c).fill = V_FILL[v]
            g.cell(row=g.max_row, column=c).font = V_FONT[v]
    g.append([])
    g.append(["MUC DO - neu phai xu ly thi gap den dau?"])
    g.cell(row=g.max_row, column=1).font = Font(bold=True, size=12)
    g.append(["Muc do", "Mau", "Y nghia"])
    for c in range(1, 4):
        g.cell(row=g.max_row, column=c).fill = HEADER_FILL
        g.cell(row=g.max_row, column=c).font = HEADER_FONT
    legend = [
        (1, "Link mat han. Phai thay bang nguon moi."),
        (2, "Link con nhung dang hong. Xu ly som."),
        (3, "Link con, gia tri thap hon ky vong hoac chua ket luan duoc."),
        (4, "Khong sai, chi ghi chu de biet."),
        (5, "Tot, khong can lam gi."),
    ]
    for sev, mean in legend:
        g.append([D.SEV_LABEL[sev], "", mean])
        for c in range(1, 4):
            g.cell(row=g.max_row, column=c).fill = FILL[sev]
            g.cell(row=g.max_row, column=c).font = FONT[sev]
    g.append([])
    g.append(["Luu y", "", "Muc do da duoc dieu chinh theo 'priority' cua tung tier "
                          "trong file config: tier uu tien 1 bi nang muc, tier uu tien "
                          "tu 4 tro len duoc ha muc."])
    g.append(["Luu y", "", "Cot 'Khop tang' cho biet link thuc su roi vao tang nao "
                          "(money / tier 2 / tier 3...). Lech so voi dich dau tien khai "
                          "trong 'targets' cua tier thi mang ma TRO_SAI_TANG - link van "
                          "con, chi la so do tang tren giay khong khop thuc te."])
    for col, w in zip("ABC", (12, 8, 96)):
        g.column_dimensions[col].width = w

    wb.save(path)
    return path
